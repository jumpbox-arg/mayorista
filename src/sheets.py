# -*- coding: utf-8 -*-
"""
Capa de acceso al CRM (Google Sheet) — lectura y escritura seguras.

Dos formas de autenticar (ambas soportadas):
  - Opción A (OAuth): login con TU usuario de Google. No usa cuenta de
    servicio, así que no la frena la política de la org. Ver docs/SETUP_OAUTH.md.
  - Opción 2 (cuenta de servicio): credencial JSON. Ver docs/SETUP_CREDENCIAL.md.

REGLAS DE SEGURIDAD (no negociables, están en el código a propósito):
  - Nunca borra filas ni columnas.
  - En el enriquecimiento sólo escribe las columnas EMAIL y NOTAS.
  - En el scraping sólo AGREGA filas nuevas al final.
  - Toda escritura pasa por métodos explícitos; el modo lectura es el default.

Incluye además un modo CSV (sin credenciales) como plan B.
"""
import csv
from pathlib import Path


class CRM:
    """Envuelve un worksheet de gspread y mapea columnas lógicas <-> reales."""

    def __init__(self, worksheet, columnas: dict, fila_encabezado: int = 1):
        self.ws = worksheet
        self.columnas = columnas                    # logico -> nombre real
        self.fila_encabezado = fila_encabezado
        self._encabezados = worksheet.row_values(fila_encabezado)
        # nombre real de columna -> índice de columna (1-based) en el Sheet
        self._col_idx = {h: i + 1 for i, h in enumerate(self._encabezados)}

    SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

    # ---- factory: cuenta de servicio (Opción 2) --------------------------
    @classmethod
    def abrir(cls, cfg, ruta_credencial: str) -> "CRM":
        import gspread
        from google.oauth2.service_account import Credentials
        creds = Credentials.from_service_account_file(ruta_credencial, scopes=cls.SCOPES)
        return cls._desde_creds(cfg, creds)

    # ---- factory: login con tu propio usuario (Opción A, OAuth) ----------
    @classmethod
    def abrir_oauth(cls, cfg, client_secret: str,
                    token_cache: str = "credenciales/token.json") -> "CRM":
        """
        Entra al Sheet como TU usuario de Google (dueño del Sheet).
        No usa cuenta de servicio, así que no lo frena la política de la org.
        La primera vez abre el navegador para que autorices; después guarda un
        token y ya no vuelve a pedir login.
        """
        import gspread
        from google.oauth2.credentials import Credentials as UserCreds
        from google_auth_oauthlib.flow import InstalledAppFlow
        from google.auth.transport.requests import Request

        creds = None
        if Path(token_cache).exists():
            creds = UserCreds.from_authorized_user_file(token_cache, cls.SCOPES)
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                flow = InstalledAppFlow.from_client_secrets_file(client_secret, cls.SCOPES)
                creds = flow.run_local_server(port=0)  # abre el navegador una vez
            Path(token_cache).write_text(creds.to_json(), encoding="utf-8")
        return cls._desde_creds(cfg, creds)

    @classmethod
    def _desde_creds(cls, cfg, creds) -> "CRM":
        import gspread
        gc = gspread.authorize(creds)
        sh = gc.open_by_key(cfg.sheet["id"])
        ws = sh.worksheet(cfg.sheet.get("hoja")) if cfg.sheet.get("hoja") else sh.sheet1
        return cls(ws, cfg.columnas, cfg.sheet.get("fila_encabezado", 1))

    # ---- lectura ----------------------------------------------------------
    def leer_filas(self):
        """Devuelve lista de dicts {campo_logico: valor} con nº de fila real."""
        valores = self.ws.get_all_values()
        cabecera = valores[self.fila_encabezado - 1]
        inv = {v: k for k, v in self.columnas.items()}  # real -> logico
        filas = []
        for n, fila in enumerate(valores[self.fila_encabezado:], start=self.fila_encabezado + 1):
            registro = {"_fila": n}
            for i, celda in enumerate(fila):
                real = cabecera[i] if i < len(cabecera) else None
                if real in inv:
                    registro[inv[real]] = celda
            filas.append(registro)
        return filas

    # ---- escritura (Trabajo 1) -------------------------------------------
    def escribir_email(self, fila_real: int, email: str = None, nota: str = None):
        """Escribe SOLO las celdas EMAIL y/o NOTAS de una fila. Nada más."""
        peticiones = []
        if email is not None:
            col = self._col_idx[self.columnas["email"]]
            peticiones.append({"range": _a1(fila_real, col), "values": [[email]]})
        if nota is not None:
            col = self._col_idx[self.columnas["notas"]]
            peticiones.append({"range": _a1(fila_real, col), "values": [[nota]]})
        # batch_update tiene firma estable en gspread 5.x y 6.x (a diferencia de
        # worksheet.update, que cambió el orden de argumentos entre versiones).
        if peticiones:
            self.ws.batch_update(peticiones, value_input_option="USER_ENTERED")

    # ---- escritura (Trabajo 2) -------------------------------------------
    def agregar_filas(self, filas_dict: list):
        """Agrega filas nuevas AL FINAL, respetando el orden de columnas real."""
        matriz = []
        for reg in filas_dict:
            fila = []
            for real in self._encabezados:
                logico = {v: k for k, v in self.columnas.items()}.get(real)
                fila.append(reg.get(logico, "") if logico else "")
            matriz.append(fila)
        if matriz:
            self.ws.append_rows(matriz, value_input_option="USER_ENTERED")


def _a1(fila: int, col: int) -> str:
    """(fila, col) 1-based -> notación A1 (ej. 8, 8 -> 'H8')."""
    letras = ""
    while col:
        col, r = divmod(col - 1, 26)
        letras = chr(65 + r) + letras
    return f"{letras}{fila}"


# --------------------------------------------------------------------------
# Modo CSV (plan B, sin credenciales): trabaja sobre un archivo local.
# --------------------------------------------------------------------------
class CRMCsv:
    def __init__(self, ruta: str, columnas: dict):
        self.ruta = Path(ruta)
        self.columnas = columnas
        with self.ruta.open(encoding="utf-8") as f:
            self._filas = list(csv.DictReader(f))
        self._campos = list(self._filas[0].keys()) if self._filas else []

    def leer_filas(self):
        inv = {v: k for k, v in self.columnas.items()}
        out = []
        for n, fila in enumerate(self._filas, start=2):
            reg = {"_fila": n}
            for real, val in fila.items():
                if real in inv:
                    reg[inv[real]] = val
            out.append(reg)
        return out

    def escribir_email(self, fila_real, email=None, nota=None):
        idx = fila_real - 2
        if email is not None:
            self._filas[idx][self.columnas["email"]] = email
        if nota is not None:
            self._filas[idx][self.columnas["notas"]] = nota

    def guardar(self, ruta_salida=None):
        destino = Path(ruta_salida or self.ruta)
        with destino.open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=self._campos)
            w.writeheader()
            w.writerows(self._filas)


# --------------------------------------------------------------------------
# CRM como archivo Excel local (.xlsx) — 100% sin Google, todo automático.
# El sistema lee y escribe el mismo archivo en tu compu. Ideal cuando la
# organización bloquea el acceso por API al Google Sheet.
# --------------------------------------------------------------------------
class CRMXlsx:
    def __init__(self, ruta: str, columnas: dict, hoja: str = None, fila_encabezado: int = 1):
        from openpyxl import load_workbook
        self.ruta = Path(ruta)
        self.columnas = columnas
        self.fila_encabezado = fila_encabezado
        self.wb = load_workbook(self.ruta)
        self.ws = self.wb[hoja] if hoja and hoja in self.wb.sheetnames else self.wb.active
        self._encabezados = [c.value for c in self.ws[fila_encabezado]]
        # nombre real de columna -> índice (1-based)
        self._col_idx = {h: i + 1 for i, h in enumerate(self._encabezados) if h}

    def leer_filas(self):
        inv = {v: k for k, v in self.columnas.items()}  # real -> logico
        filas = []
        for n in range(self.fila_encabezado + 1, self.ws.max_row + 1):
            reg = {"_fila": n}
            for real, col in self._col_idx.items():
                if real in inv:
                    val = self.ws.cell(row=n, column=col).value
                    reg[inv[real]] = "" if val is None else str(val)
            filas.append(reg)
        return filas

    def escribir_email(self, fila_real, email=None, nota=None):
        if email is not None:
            self.ws.cell(row=fila_real, column=self._col_idx[self.columnas["email"]], value=email)
        if nota is not None:
            self.ws.cell(row=fila_real, column=self._col_idx[self.columnas["notas"]], value=nota)

    def agregar_filas(self, filas_dict: list):
        inv = {v: k for k, v in self.columnas.items()}
        for reg in filas_dict:
            fila = []
            for real in self._encabezados:
                logico = inv.get(real)
                fila.append(reg.get(logico, "") if logico else "")
            self.ws.append(fila)

    def guardar(self, ruta_salida=None):
        self.wb.save(ruta_salida or self.ruta)
